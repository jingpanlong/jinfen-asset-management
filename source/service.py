from __future__ import annotations
import csv
import contextlib
import hashlib
import hmac
import json
import re
import secrets
import uuid
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from database import Database, SCHEMA, COLUMNS, INTEGER_COLUMNS
from access import AccessMixin,validate_policy,defaults,BUILTIN_USERNAME

STATUS = ['闲置', '在用', '维修中', '已报废']
ROLES = ['管理员', '操作员', '只读用户']
ASSET_FIELDS = ['id','name','category','brand','serial','quantity','owner','location','supplier_id','purchase_date','price_cents','life_months','residual_bp','warranty_date','notes']
LABELS = {'id':'编号','name':'设备名称','category':'类别','brand':'品牌型号','serial':'序列号','quantity':'数量','owner':'归属','location':'存放位置','supplier_id':'供应商编号','purchase_date':'购入日期','price_cents':'单价','life_months':'使用月数','residual_bp':'残值率','warranty_date':'保修到期','status':'状态','employee_id':'当前员工编号','due_date':'应还日期','notes':'其他备注'}

class BusinessError(ValueError):
    pass

def stamp():
    return datetime.now(timezone.utc).isoformat(timespec='seconds')

def ident():
    return str(uuid.uuid4())

def clean(value, maximum=5000):
    s = str(value or '').strip()
    if len(s) > maximum:
        raise BusinessError(f'文字长度不能超过 {maximum} 字')
    return s

def integer(value, minimum=0, maximum=1000000000000):
    try:
        d = Decimal(str(value))
        if not d.is_finite() or d != d.to_integral_value() or not minimum <= d <= maximum:
            raise ValueError()
        return int(d)
    except (InvalidOperation, ValueError, TypeError):
        raise BusinessError(f'整数必须在 {minimum}～{maximum} 之间')

def money(value):
    try:
        amount = Decimal(str(value or '0'))
        if not amount.is_finite() or amount < 0 or amount > Decimal('1000000000'):
            raise ValueError()
        return int((amount * 100).quantize(Decimal('1'), rounding=ROUND_HALF_UP))
    except (InvalidOperation, ValueError):
        raise BusinessError('金额必须是 0～10 亿之间的数字')

def day(value, optional=True):
    value = clean(value, 10)
    if not value and optional:
        return ''
    try:
        return date.fromisoformat(value).isoformat()
    except ValueError:
        raise BusinessError('日期格式必须为 YYYY-MM-DD')

def pw_hash(password):
    if len(password) < 10 or len(password) > 200 or not re.search(r'[A-Za-z]', password) or not re.search(r'\d', password):
        raise BusinessError('密码需 10～200 位，包含英文字母和数字')
    salt = secrets.token_hex(16)
    digest = hashlib.pbkdf2_hmac('sha256', password.encode(), bytes.fromhex(salt), 310000).hex()
    return 'pbkdf2_sha256$310000$' + salt + '$' + digest

def pw_check(password, stored):
    try:
        method, rounds, salt, digest = stored.split('$')
        if method != 'pbkdf2_sha256' or not 100000 <= int(rounds) <= 2000000:
            return False
        actual = hashlib.pbkdf2_hmac('sha256', password.encode(), bytes.fromhex(salt), int(rounds)).hex()
        return hmac.compare_digest(actual, digest)
    except (ValueError, TypeError):
        return False

class Service(AccessMixin):
    def __init__(self, db):
        self.db: Database = db
        self.user_id = None

    def audit(self, action, detail, actor=None):
        self.db.insert('audit', {'id':ident(),'actor_id':actor or self.user_id or 'system','action':action,'detail':clean(detail, 50000),'created_at':stamp()})

    def require(self, admin=False, write=False):
        u = self.db.one('SELECT * FROM users WHERE id=?', (self.user_id,))
        if not u or not u['active']:
            raise BusinessError('账号已停用或登录失效，请重新登录')
        epoch=self.db.one("SELECT value FROM meta WHERE key='auth_epoch'")
        if getattr(self,'session_password',None) and (self.session_password!=u['password'] or self.session_epoch!=(epoch['value'] if epoch else '')):
            raise BusinessError('账号凭据已重置或数据库已恢复，请重新登录')
        p=self.profile(u)
        if p['expires'] and datetime.fromisoformat(p['expires'])<=datetime.now(timezone.utc):
            raise BusinessError('账号授权已到期，请联系管理员')
        if getattr(self,'session_access',None) and self.session_access!=self.access_signature(u):
            raise BusinessError('权限已变化，请重新登录')
        writable=any(not k.endswith('.read') and k not in ['data.export','data.print'] for k in p['permissions'])
        if (admin and u['role'] != '管理员') or (write and not writable):
            raise BusinessError('当前账号没有此操作权限')
        return u

    def register(self, username, display_name, password):
        """Self-registration: always an ordinary read-only account, never an administrator."""
        username = clean(username, 40).lower()
        name = clean(display_name, 80)
        if not re.fullmatch(r'[a-z][a-z0-9_.-]{2,39}', username) or not name:
            raise BusinessError('用户名需 3～40 位，以英文字母开头，可含数字、下划线、点和短横线；姓名必填')
        if username == BUILTIN_USERNAME:
            raise BusinessError('该用户名为软件内置管理员保留，请换一个用户名')
        hashed = pw_hash(password)
        with self.db.transaction():
            self.ensure_builtin_owner()
            if self.db.one('SELECT id FROM users WHERE username=?', (username,)):
                raise BusinessError('用户名已存在')
            # Registration never yields administration: the built-in account is
            # the one and only way to grant it, and the account starts usable
            # with read-only rights.
            user = {'id':ident(),'username':username,'display_name':name,'password':hashed,'role':'只读用户','active':1,'failures':0,'locked_until':'','created_at':stamp()}
            self.db.insert('users', user)
            self.audit('注册', f'{username}：只读账号（管理员只能由内置管理员账号授予）', user['id'])
        return False

    def login(self, username, password):
        error = None
        with self.db.transaction():
            # Repair the built-in administrator *before* the password check:
            # otherwise a tampered hash would lock the factory account out.
            self.ensure_builtin_owner()
            u = self.db.one('SELECT * FROM users WHERE username=?', (clean(username, 40).lower(),))
            if not u:
                # A dummy expensive hash reduces username timing disclosure.
                pw_check(password, 'pbkdf2_sha256$310000$'+'00'*16+'$'+'00'*32)
                error = '用户名或密码错误'
            elif u['locked_until'] and u['locked_until'] > stamp():
                error = '密码错误次数过多，请 15 分钟后再试，或联系管理员解锁'
            elif not pw_check(password, u['password']):
                failures = u['failures'] + 1 if u['locked_until'] > stamp() or not u['locked_until'] else 1
                locked = (datetime.now(timezone.utc)+timedelta(minutes=15)).isoformat(timespec='seconds') if failures >= 5 else ''
                self.db.update('users',u['id'],{'failures':failures,'locked_until':locked})
                self.audit('登录失败',u['username'],u['id'])
                error = '用户名或密码错误'
            elif not u['active']:
                error = '账号待管理员审核或已停用'
            else:
                self.user_id = u['id']
                self.session_access = self.access_signature(u)
                self.session_password = u['password']
                epoch=self.db.one("SELECT value FROM meta WHERE key='auth_epoch'")
                self.session_epoch=epoch['value'] if epoch else ''
                self.db.update('users',u['id'],{'failures':0,'locked_until':''})
                self.audit('登录',u['username'])
        if error:
            raise BusinessError(error)
        return self.require()

    def logout(self):
        if self.user_id:
            with self.db.transaction():
                self.audit('退出登录','用户退出系统')
        self.user_id = None
        self.session_access = None

    def change_password(self, old, new):
        hashed = pw_hash(new)
        with self.db.transaction():
            u = self.require()
            if str(u['id']) == self.owner_id():
                raise BusinessError('内置管理员是出厂固定账号，口令不能在软件内修改')
            if not pw_check(old, u['password']):
                raise BusinessError('原密码错误')
            self.db.update('users',u['id'],{'password':hashed,'failures':0,'locked_until':''})
            self.session_password=hashed
            self.audit('修改密码','当前用户修改密码')

    def manage_user(self, id, role, active, new_password=''):
        hashed = pw_hash(new_password) if new_password else None
        if role not in ROLES:
            raise BusinessError('角色无效')
        with self.db.transaction():
            me = self.require(admin=True)
            owner = self.is_owner(me)
            u = self.db.one('SELECT * FROM users WHERE id=?',(id,))
            if not u:
                raise BusinessError('账号不存在')
            if str(u['id'])==self.owner_id():
                raise BusinessError('内置管理员是出厂固定账号，其角色、启用状态和口令均不可改动')
            if not owner:
                # A granted administrator runs ordinary accounts only; creating or
                # touching administrators is the built-in account's decision alone.
                if u['role']=='管理员' or role=='管理员':
                    raise BusinessError('只有内置管理员账号可以新建或调整管理员，请交由内置管理员操作')
            if id == self.user_id and (not active or role != '管理员'):
                raise BusinessError('不能停用或降低当前管理员自己的权限')
            admins = self.db.all("SELECT id FROM users WHERE role='管理员' AND active=1")
            if u['role']=='管理员' and u['active'] and len(admins)<=1 and (not active or role!='管理员'):
                raise BusinessError('至少保留一个启用的管理员')
            fields = {'role':role,'active':int(bool(active)),'failures':0,'locked_until':''}
            if hashed:
                fields['password']=hashed
            self.db.update('users',id,fields)
            self.audit('账号管理', f"{u['username']}：{role} / {'启用' if active else '停用'}"+(' / 重置密码' if hashed else ''))

    def list(self, table):
        if table in ['users','audit']:self.permit('security.manage' if table=='users' else 'audit.read')
        return self.scoped_rows(table)

    def catalog(self, table, values, editing=False):
        if table not in ['departments','employees','suppliers']:
            raise BusinessError('基础资料类别无效')
        self.permit(table+'.write');self.global_scope()
        data = {k:clean(values.get(k,''),3000) for k in COLUMNS[table]}
        data['id']=clean(data['id'],100)
        if not data['id'] or not data['name']:
            raise BusinessError('编号和名称必填')
        if table=='employees':
            data['active']=int(bool(values.get('active',1)))
        with (contextlib.nullcontext() if self.db.inside else self.db.transaction()):
            self.require(write=True)
            exists=self.db.one(f'SELECT * FROM {table} WHERE id=?',(data['id'],))
            if bool(exists)!=bool(editing):
                raise BusinessError('编号已存在或记录已不存在，请刷新')
            if table in ['departments','suppliers']:
                duplicate=self.db.one(f'SELECT id FROM {table} WHERE name=? AND id<>?',(data['name'],data['id']))
                if duplicate:
                    raise BusinessError('名称已存在')
            if table=='employees':
                if data['department_id'] and not self.db.one('SELECT id FROM departments WHERE id=?',(data['department_id'],)):
                    raise BusinessError('部门不存在')
                if not data['active'] and self.db.one('SELECT id FROM assets WHERE employee_id=?',(data['id'],)):
                    raise BusinessError('该员工还有在用资产，请先归还或调拨')
            if exists:
                self.db.update(table,data['id'],{k:v for k,v in data.items() if k!='id'})
            else:
                self.db.insert(table,data)
            self.audit('基础资料编辑' if exists else '基础资料新增',table+' / '+data['id']+' / '+data['name'])

    def asset_data(self, values):
        a = {k:clean(values.get(k,'')) for k in ASSET_FIELDS}
        for k in ['id','name','category']:
            if not a[k] or len(a[k])>300:
                raise BusinessError('编号、设备名称、类别必填，且最多 300 字')
        a['quantity']=integer(values.get('quantity',1),1,1000000)
        a['price_cents']=integer(values.get('price_cents',0))
        a['life_months']=integer(values.get('life_months',36),1,1200)
        a['residual_bp']=integer(values.get('residual_bp',500),0,10000)
        a['purchase_date']=day(a['purchase_date'])
        a['warranty_date']=day(a['warranty_date'])
        if a['purchase_date'] and a['warranty_date'] and a['warranty_date']<a['purchase_date']:
            raise BusinessError('保修到期不能早于购入日期')
        return a

    def event(self, a, action, detail):
        self.db.insert('events',{'id':ident(),'asset_id':a['id'],'action':action,'actor_id':self.user_id,'detail':clean(detail,50000),'created_at':stamp()})
        self.audit(action,'#'+a['id']+' '+detail)

    def _save_asset(self, values, revision=None):
        self.permit('assets.create' if revision is None else 'assets.edit')
        if revision is not None:self.check_asset(values.get('id'))
        c=self.profile()
        if c['scope']=='self':raise BusinessError('本人领用范围不允许新增或编辑台账')
        if c['scope']=='locations' and values.get('location') not in c['locations']:raise BusinessError('目标位置不在授权范围内')
        a = self.asset_data(values)
        if a['supplier_id'] and not self.db.one('SELECT id FROM suppliers WHERE id=?',(a['supplier_id'],)):
            raise BusinessError('供应商不存在')
        old=self.db.one('SELECT * FROM assets WHERE id=?',(a['id'],))
        if revision is None and old:
            raise BusinessError('资产编号已存在：'+a['id'])
        if revision is not None and (not old or old['revision'] != revision):
            raise BusinessError('资产已被其他操作修改，请刷新后重试')
        if old:
            if old['status']=='已报废':
                raise BusinessError('报废资产仅可查看历史，不能编辑')
            changes={k:v for k,v in a.items() if k!='id' and v!=old[k]}
            if not changes:
                return old
            detail='；'.join(f'{LABELS[k]}：{old[k]} → {v}' for k,v in changes.items())
            changes.update(revision=revision+1,updated_at=stamp())
            self.db.update('assets',a['id'],changes)
            self.event(a,'编辑资产',detail)
        else:
            a.update(status='闲置',employee_id='',due_date='',revision=1,created_at=stamp(),updated_at=stamp())
            self.db.insert('assets',a)
            self.event(a,'资产入库',a['name']+' / '+str(a['quantity'])+' 件')
        return self.db.one('SELECT * FROM assets WHERE id=?',(a['id'],))

    def save_asset(self, values, revision=None):
        with self.db.transaction():
            self.require(write=True)
            return self._save_asset(values,revision)

    def flow(self, id, revision, action, employee_id='', due_date='', location='', note='', vendor='', cost='0'):
        self.permit({'领用':'assets.issue','归还':'assets.issue','调拨':'assets.transfer','移动位置':'assets.transfer','送修':'assets.repair','维修完成':'assets.repair','报废':'assets.scrap'}.get(action,'invalid'))
        self.check_asset(id)
        c=self.profile()
        if action=='移动位置' and c['scope']=='locations' and location not in c['locations']:raise BusinessError('目标位置不在授权范围内')
        if action in ['领用','调拨'] and c['scope']=='self' and employee_id!=c['employee_id']:raise BusinessError('不能向授权范围外员工调拨资产')
        allowed={'闲置':['领用','送修','报废','移动位置'],'在用':['归还','调拨','送修','报废','移动位置'],'维修中':['维修完成','报废'],'已报废':[]}
        with (contextlib.nullcontext() if self.db.inside else self.db.transaction()):
            self.require(write=True)
            a=self.db.one('SELECT * FROM assets WHERE id=?',(id,))
            if not a or a['revision']!=revision:
                raise BusinessError('资产已变化，请刷新后重试')
            if action not in allowed[a['status']]:
                raise BusinessError('此状态不支持所选操作')
            note=clean(note)
            if action in ['送修','报废','维修完成'] and not note:
                raise BusinessError('此操作必须填写原因 / 处理结果')
            fields={'revision':revision+1,'updated_at':stamp()}
            detail='原状态：'+a['status']+'；原使用员工：'+(a['employee_id'] or '无')
            if action in ['领用','调拨']:
                emp=self.db.one('SELECT * FROM employees WHERE id=?',(employee_id,))
                if not emp or not emp['active']:
                    raise BusinessError('请选择启用的员工')
                due_date=day(due_date)
                if due_date and due_date<date.today().isoformat():
                    raise BusinessError('应还日期不能早于今天')
                if action=='调拨' and employee_id==a['employee_id']:
                    raise BusinessError('调拨员工与当前使用员工相同')
                fields.update(status='在用',employee_id=employee_id,due_date=due_date)
                detail+='；新员工：'+emp['name']+' / '+employee_id+'；应还：'+(due_date or '未设定')
            elif action=='移动位置':
                location=clean(location,300)
                if not location or location==a['location']:
                    raise BusinessError('请填写与原位置不同的新位置')
                fields['location']=location
                detail+='；位置：'+a['location']+' → '+location
            else:
                fields.update(status={'归还':'闲置','送修':'维修中','维修完成':'闲置','报废':'已报废'}[action],employee_id='',due_date='')
            if action=='送修':
                self.db.insert('repairs',{'id':ident(),'asset_id':id,'problem':note,'vendor':clean(vendor),'cost_cents':0,'result':'','status':'维修中','opened_at':stamp(),'closed_at':''})
            elif action in ['维修完成','报废'] and a['status']=='维修中':
                repair=self.db.one("SELECT * FROM repairs WHERE asset_id=? AND status='维修中'",(id,))
                if not repair:
                    raise BusinessError('维修工单缺失，请联系管理员检查数据')
                self.db.update('repairs',repair['id'],{'status':'已完成' if action=='维修完成' else '报废结束','result':note,'cost_cents':money(cost),'closed_at':stamp()})
                detail+='；维修费用：'+format(Decimal(money(cost))/100,'.2f')
            self.db.update('assets',id,fields)
            self.event(a,action,detail+'；说明：'+note)

    def new_purchase(self, values):
        self.permit('purchases.create')
        name=clean(values.get('name'),300); reason=clean(values.get('reason'))
        if not name or not reason:
            raise BusinessError('采购设备名称和理由必填')
        with self.db.transaction():
            self.require(write=True)
            id=ident()
            self.db.insert('purchases',{'id':id,'name':name,'category':clean(values.get('category') or '其他',300),'quantity':integer(values.get('quantity',1),1,1000000),'budget_cents':money(values.get('budget','0')),'reason':reason,'requester_id':self.user_id,'status':'待审批','approver_id':'','comment':'','asset_id':'','created_at':stamp(),'updated_at':stamp()})
            self.audit('采购申请',name+' / '+id)
            return id

    def approve_purchase(self, id, approved, comment):
        self.permit('purchases.approve');self.check_row('purchases',id)
        with self.db.transaction():
            self.require()
            p=self.db.one('SELECT * FROM purchases WHERE id=?',(id,))
            if not p or p['status']!='待审批':
                raise BusinessError('采购申请已处理，请刷新')
            if p['requester_id']==self.user_id:
                raise BusinessError('申请人与审批人不能是同一账号')
            comment=clean(comment)
            if not approved and not comment:
                raise BusinessError('拒绝采购必须填写理由')
            self.db.update('purchases',id,{'status':'已批准' if approved else '已拒绝','approver_id':self.user_id,'comment':comment,'updated_at':stamp()})
            self.audit('采购审批',id+' / '+('批准' if approved else '拒绝')+' / '+comment)

    def receive_purchase(self, id, values):
        self.permit('purchases.receive');self.check_row('purchases',id)
        with self.db.transaction():
            self.require(write=True)
            p=self.db.one('SELECT * FROM purchases WHERE id=?',(id,))
            if not p or p['status']!='已批准':
                raise BusinessError('只有已批准的采购可以入库')
            a=self._save_asset(values)
            if a['quantity']!=p['quantity']:
                raise BusinessError('入库数量必须与采购申请一致')
            self.db.update('purchases',id,{'status':'已入库','asset_id':a['id'],'updated_at':stamp()})
            self.audit('采购入库',id+' → '+a['id'])

    def new_count(self, name, location=''):
        self.permit('counts.write')
        name=clean(name,300); location=clean(location,300)
        if not name:
            raise BusinessError('盘点名称必填')
        with self.db.transaction():
            self.require(write=True)
            assets=self.db.all("SELECT * FROM assets WHERE status<>'已报废'"+(' AND location=?' if location else ''),(location,) if location else ())
            assets=[a for a in assets if self.asset_visible(a)]
            if not assets:
                raise BusinessError('所选范围没有可盘点资产')
            id=ident()
            self.db.insert('counts',{'id':id,'name':name,'location':location,'status':'进行中','creator_id':self.user_id,'created_at':stamp(),'closed_at':''})
            for a in assets:
                self.db.insert('count_items',{'id':ident(),'count_id':id,'asset_id':a['id'],'expected_quantity':a['quantity'],'expected_status':a['status'],'actual_quantity':-1,'result':'待盘点','notes':'','checked_at':'','checker_id':''})
            self.audit('建立盘点',name+' / '+str(len(assets))+' 条')
            return id

    def check_item(self, id, actual, notes):
        self.permit('counts.write');self.check_row('count_items',id)
        with self.db.transaction():
            self.require(write=True)
            item=self.db.one('SELECT * FROM count_items WHERE id=?',(id,))
            if not item:
                raise BusinessError('盘点项目不存在')
            count=self.db.one('SELECT * FROM counts WHERE id=?',(item['count_id'],))
            if count['status']!='进行中':
                raise BusinessError('盘点已结束，不能修改')
            actual=integer(actual,0,1000000)
            self.db.update('count_items',id,{'actual_quantity':actual,'result':'一致' if actual==item['expected_quantity'] else '差异','notes':clean(notes),'checked_at':stamp(),'checker_id':self.user_id})
            self.audit('盘点登记',item['asset_id']+' / 实盘 '+str(actual))

    def close_count(self,id):
        self.permit('counts.write');self.check_row('counts',id)
        for item in self.db.all('SELECT asset_id FROM count_items WHERE count_id=?',(id,)):self.check_asset(item['asset_id'])
        with self.db.transaction():
            self.require(write=True)
            c=self.db.one('SELECT * FROM counts WHERE id=?',(id,))
            if not c or c['status']!='进行中':
                raise BusinessError('盘点已结束或不存在')
            if self.db.one("SELECT id FROM count_items WHERE count_id=? AND result='待盘点'",(id,)):
                raise BusinessError('还有未盘点项目，请全部登记后再结束')
            self.db.update('counts',id,{'status':'已完成','closed_at':stamp()})
            self.audit('结束盘点',id+'；盘点差异保留，资产数量不自动修改')

    # ---- Deletion and editing: assets may only be removed while nothing references them.
    def delete_asset(self, id):
        self.permit('assets.delete');self.check_asset(id)
        with self.db.transaction():
            self.require(write=True)
            a=self.check_asset(id)
            blocked=[]
            if a['employee_id']:blocked.append('当前领用关系')
            for table,label in [('repairs','维修工单'),('attachments','附件 / 图片'),('count_items','盘点明细'),('purchases','采购入库记录'),('tasks','关联待办')]:
                if self.db.one(f'SELECT id FROM {table} WHERE asset_id=?',(id,)):blocked.append(label)
            flow=sorted({e['action'] for e in self.db.all('SELECT action FROM events WHERE asset_id=?',(id,)) if e['action']!='资产入库'})
            if flow:blocked.append('流转历史（'+'/'.join(flow)+'）')
            if blocked:
                raise BusinessError('该资产已有'+ '、'.join(blocked) +'，不能删除。请改用「报废」保留完整历史，或先在对应模块中清理引用。')
            self.db.execute("DELETE FROM events WHERE asset_id=? AND action='资产入库'",(id,))
            self.db.execute('DELETE FROM assets WHERE id=?',(id,))
            self.audit('删除资产',id+' / '+a['name'])

    def delete_catalog(self, table, id):
        if table not in ['departments','employees','suppliers']:raise BusinessError('该类型资料不支持删除')
        self.permit(table+'.delete');self.global_scope();self.check_row(table,id)
        with self.db.transaction():
            self.require(write=True)
            row=self.db.one(f'SELECT * FROM {table} WHERE id=?',(id,))
            if not row:raise BusinessError('资料不存在')
            if table=='departments':used,label=self.db.one('SELECT COUNT(*) AS n FROM employees WHERE department_id=?',(id,))['n'],'员工档案'
            elif table=='employees':used,label=self.db.one('SELECT COUNT(*) AS n FROM assets WHERE employee_id=?',(id,))['n'],'资产'
            else:used,label=self.db.one('SELECT COUNT(*) AS n FROM assets WHERE supplier_id=?',(id,))['n'],'资产'
            if used:raise BusinessError(f'仍有 {used} 条{label}引用该资料，请先调整归属后再删除')
            self.db.execute(f'DELETE FROM {table} WHERE id=?',(id,))
            self.audit('删除资料',table+' / '+id+' / '+str(row.get('name','')))

    def delete_user(self, id):
        self.permit('users.delete')
        with self.db.transaction():
            me=self.require(write=True)
            u=self.db.one('SELECT * FROM users WHERE id=?',(id,))
            if not u:raise BusinessError('账号不存在')
            if id==self.user_id:raise BusinessError('不能删除当前正在登录的账号')
            if str(u['id'])==self.owner_id():raise BusinessError('内置管理员是出厂固定账号，不能删除')
            if u['role']=='管理员' and not self.is_owner(me):raise BusinessError('只有内置管理员账号可以删除管理员账号')
            if u['role']=='管理员' and self.db.one("SELECT COUNT(*) AS n FROM users WHERE role='管理员'")['n']<=1:
                raise BusinessError('至少需要保留一名管理员账号')
            self.db.execute('DELETE FROM users WHERE id=?',(id,))
            policy=self.policy();policy['users'].pop(id,None);self._write_policy(policy)
            self.audit('删除账号',u['username']+' / '+u['display_name'])

    def update_purchase(self, id, values):
        self.permit('purchases.edit');self.check_row('purchases',id)
        name=clean(values.get('name'),300);reason=clean(values.get('reason'))
        if not name or not reason:raise BusinessError('采购设备名称和理由必填')
        with self.db.transaction():
            self.require(write=True)
            p=self.db.one('SELECT * FROM purchases WHERE id=?',(id,))
            if not p:raise BusinessError('采购申请不存在')
            if p['status']!='待审批':raise BusinessError('只有「待审批」的申请可以修改内容，已处理的申请请重新提交')
            self.db.update('purchases',id,{'name':name,'category':clean(values.get('category') or '其他',300),'quantity':integer(values.get('quantity',1),1,1000000),'budget_cents':money(values.get('budget','0')),'reason':reason,'updated_at':stamp()})
            self.audit('修改采购申请',name+' / '+id)

    def delete_purchase(self, id):
        self.permit('purchases.delete');self.check_row('purchases',id)
        with self.db.transaction():
            self.require(write=True)
            p=self.db.one('SELECT * FROM purchases WHERE id=?',(id,))
            if not p:raise BusinessError('采购申请不存在')
            if p['status'] not in ['待审批','已拒绝']:
                raise BusinessError('「'+p['status']+'」的申请保留审批与入库记录，不能删除')
            self.db.execute('DELETE FROM purchases WHERE id=?',(id,))
            self.audit('删除采购申请',p['name']+' / '+id)

    def delete_count(self, id):
        self.permit('counts.delete');self.check_row('counts',id)
        with self.db.transaction():
            self.require(write=True)
            c=self.db.one('SELECT * FROM counts WHERE id=?',(id,))
            if not c:raise BusinessError('盘点记录不存在')
            items=self.db.one('SELECT COUNT(*) AS n FROM count_items WHERE count_id=?',(id,))['n']
            self.db.execute('DELETE FROM count_items WHERE count_id=?',(id,))
            self.db.execute('DELETE FROM counts WHERE id=?',(id,))
            self.audit('删除盘点',c['name']+' / 明细 '+str(items)+' 条一并删除')

    LOG_LABELS={'events':'流转记录','audit':'操作审计','import_batches':'导入履历'}

    def delete_log(self, table, id):
        labels=self.LOG_LABELS
        if table not in labels:raise BusinessError('该类型记录不支持删除')
        self.permit('audit.delete')
        with self.db.transaction():
            self.require(write=True)
            self.check_row(table,id)
            row=self.db.one(f'SELECT * FROM {table} WHERE id=?',(id,))
            if not row:raise BusinessError(labels[table]+'不存在')
            self.db.execute(f'DELETE FROM {table} WHERE id=?',(id,))
            self.audit('删除'+labels[table],id)

    def depreciation(self, a, today=None):
        total=a['price_cents']*a['quantity']
        if a['status']=='已报废':
            return total, total, 0
        if not a['purchase_date']:
            return total, 0, total
        start=date.fromisoformat(a['purchase_date']); end=today or date.today()
        elapsed=max(0,(end.year-start.year)*12+end.month-start.month-(end.day<start.day))
        months=min(a['life_months'],elapsed)
        depreciable=Decimal(total)*(10000-a['residual_bp'])/10000
        dep=int((depreciable*months/a['life_months']).quantize(Decimal('1'),rounding=ROUND_HALF_UP))
        return total,dep,total-dep

    def backup(self, path):
        self.permit('data.backup');self.global_scope()
        with self.db.transaction():
            self.require()
            data={'format':'qiyun-desktop','version':3,'created_at':stamp(),'tables':{t:self.db.rows(t) for t in SCHEMA if t!='meta'}}
            data['access_policy']=self.policy()
            self.audit('创建备份','导出完整数据库备份')
        Path(path).write_text(json.dumps(data,ensure_ascii=False,indent=2),'utf-8')

    def validate_backup(self, data):
        if data.get('format')=='qiyun-desktop' and data.get('version')==2:
            old_tables=set(SCHEMA)-{'meta','attachments','tasks','import_batches'}
            if set(data.get('tables',{}))!=old_tables:
                raise BusinessError('旧版本备份字段不完整')
            data=json.loads(json.dumps(data));data['version']=3
            for t in ['attachments','tasks','import_batches']:data['tables'][t]=[]
        if data.get('format')!='qiyun-desktop' or data.get('version')!=3 or set(data.get('tables',{})) != set(SCHEMA)-{'meta'}:
            raise BusinessError('不是本软件的有效完整备份')
        tables=data['tables']; ids={}
        for table,rows in tables.items():
            if not isinstance(rows,list) or len(rows)>200000:
                raise BusinessError('备份数据格式或大小无效')
            ids[table]=set()
            for r in rows:
                if not isinstance(r,dict) or set(r)!=set(COLUMNS[table]):
                    raise BusinessError('备份字段无效：'+table)
                for k,v in r.items():
                    if k in INTEGER_COLUMNS[table]:
                        if type(v)!=int or abs(v)>10**15:
                            raise BusinessError('整数数据无效')
                    elif not isinstance(v,str) or len(v)>(15*1024*1024 if table=='attachments' and k=='content_b64' else 50000):
                        raise BusinessError('文本数据无效')
                if not r['id'] or r['id'] in ids[table]:
                    raise BusinessError('备份编号为空或重复')
                ids[table].add(r['id'])
        validate_policy(data.get('access_policy',defaults()),ids['users'])
        for config in data.get('access_policy',defaults())['users'].values():
            if config['employee_id'] and config['employee_id'] not in ids['employees']:raise BusinessError('权限配置关联的员工不存在')
        if not any(u['role']=='管理员' and u['active']==1 for u in tables['users']):
            raise BusinessError('备份必须包含启用的管理员')
        usernames=set()
        for u in tables['users']:
            if u['role'] not in ROLES or u['active'] not in [0,1] or u['username'] in usernames or not re.fullmatch(r'[a-z][a-z0-9_.-]{2,39}',u['username']) or not re.fullmatch(r'pbkdf2_sha256\$(?:[1-9][0-9]{5,6})\$[a-f0-9]{32}\$[a-f0-9]{64}',u['password']):
                raise BusinessError('备份账号数据无效')
            usernames.add(u['username'])
        def ref(table,key,target,empty=True):
            for row in tables[table]:
                if not (empty and row[key]=='') and row[key] not in ids[target]:
                    raise BusinessError('备份引用不存在：'+table+'.'+key)
        ref('employees','department_id','departments');ref('assets','supplier_id','suppliers');ref('assets','employee_id','employees');ref('events','asset_id','assets',False);ref('repairs','asset_id','assets',False);ref('purchases','requester_id','users',False);ref('purchases','approver_id','users');ref('purchases','asset_id','assets');ref('counts','creator_id','users',False);ref('count_items','count_id','counts',False);ref('count_items','asset_id','assets',False)
        ref('attachments','asset_id','assets',False);ref('tasks','asset_id','assets');ref('tasks','creator_id','users',False)
        import base64
        for item in tables['attachments']:
            try:raw=base64.b64decode(item['content_b64'],validate=True)
            except Exception:raise BusinessError('附件编码无效')
            if len(raw)!=item['size_bytes'] or len(raw)>10*1024*1024 or hashlib.sha256(raw).hexdigest()!=item['sha256']:
                raise BusinessError('附件大小或校验值无效')
            if Path(item['filename']).name!=item['filename']:
                raise BusinessError('附件文件名无效')
        if sum(a['size_bytes'] for a in tables['attachments'])>30*1024*1024:raise BusinessError('附件总量超过 30MB')
        for task in tables['tasks']:
            if not task['title'].strip() or task['status'] not in ['待处理','已完成'] or task['priority'] not in ['普通','重要','紧急']:raise BusinessError('待办数据无效')
            day(task['due_date'])
        for a in tables['assets']:
            self.asset_data(a)
            if a['status'] not in STATUS or (a['status']=='在用')!=bool(a['employee_id']) or a['revision']<1:
                raise BusinessError('备份资产状态无效')
            day(a['due_date'])
        for r in tables['repairs']:
            if r['status'] not in ['维修中','已完成','报废结束'] or r['cost_cents']<0:
                raise BusinessError('备份维修数据无效')
        for p in tables['purchases']:
            if p['status'] not in ['待审批','已批准','已拒绝','已入库'] or p['quantity']<1 or p['budget_cents']<0:
                raise BusinessError('备份采购数据无效')
        for c in tables['counts']:
            if c['status'] not in ['进行中','已完成']:
                raise BusinessError('备份盘点状态无效')
        for i in tables['count_items']:
            if i['result'] not in ['待盘点','一致','差异'] or i['expected_quantity']<1 or i['actual_quantity']< -1:
                raise BusinessError('备份盘点数据无效')
        for a in tables['assets']:
            open_repairs=[r for r in tables['repairs'] if r['asset_id']==a['id'] and r['status']=='维修中']
            if len(open_repairs)!=int(a['status']=='维修中'):
                raise BusinessError('备份维修工单与资产状态不一致')
        return data

    def restore(self,path):
        self.permit('data.restore');self.global_scope()
        if Path(path).stat().st_size>100*1024*1024:
            raise BusinessError('备份文件超过 100MB')
        data=self.validate_backup(json.loads(Path(path).read_text('utf-8')))
        with self.db.transaction():
            self.require()
            for t in reversed(list(data['tables'])):
                self.db.execute(f'DELETE FROM {t}')
            for t,rows in data['tables'].items():
                self.db.insert_many(t,rows)
            self._write_policy(data.get('access_policy',defaults()))
            self.db.execute("INSERT INTO meta (key,value) VALUES ('auth_epoch',?) ON CONFLICT (key) DO UPDATE SET value=excluded.value",(ident(),))
            self.audit('恢复备份','完整数据库恢复，所有用户需重新登录',actor='system')
        self.user_id=None
        self.session_access=None

    def import_assets(self,path):
        self.permit('data.import')
        path=Path(path)
        if path.stat().st_size>30*1024*1024:
            raise BusinessError('导入文件超过 30MB')
        browser=None
        if path.suffix.lower()=='.json':
            browser=json.loads(path.read_text('utf-8-sig'))
            if browser.get('version')!=1 or not isinstance(browser.get('assets'),list):
                raise BusinessError('这里只支持上一版网页系统的 JSON，完整桌面备份请用恢复功能')
            records=[{'编号':a.get('id'),'设备名称':a.get('name'),'数量':a.get('quantity'),'类别':a.get('category'),'品牌型号':a.get('brand'),'归属':a.get('owner'),'其他备注':a.get('notes','')+'\n网页原状态：'+str(a.get('status',''))+'；原使用人：'+str(a.get('user','')),'__history':a.get('history',[])} for a in browser['assets']]
        elif path.suffix.lower()=='.xlsx':
            from openpyxl import load_workbook
            book=load_workbook(path,read_only=True,data_only=True)
            try:
                rows=iter(book.active.iter_rows(values_only=True)); header=[clean(v) for v in next(rows)]
                if len(set(header))!=len(header):
                    raise BusinessError('Excel 存在重复表头')
                records=[dict(zip(header,row)) for row in rows if any(v is not None for v in row)]
            finally:
                book.close()
        else:
            with path.open(encoding='utf-8-sig',newline='') as f:
                reader=csv.DictReader(f)
                if not reader.fieldnames or len(set(reader.fieldnames))!=len(reader.fieldnames):
                    raise BusinessError('CSV 缺少表头或存在重复表头')
                records=list(reader)
        if not records or len(records)>50000:
            raise BusinessError('导入记录需 1～50000 条')
        with self.db.transaction():
            self.require(write=True)
            for i,row in enumerate(records,2):
                try:
                    values={k:row.get(label,'') for k,label in LABELS.items() if k in ASSET_FIELDS}
                    values['quantity']=row.get('数量',1)
                    values['price_cents']=money(row.get('单价',0))
                    values['life_months']=row.get('使用月数',36) or 36
                    values['residual_bp']=int(money(row.get('残值率',5) or 0)) # percentage to basis points
                    for key in ['purchase_date','warranty_date']:
                        if isinstance(values[key],(datetime,date)):
                            values[key]=values[key].date().isoformat() if isinstance(values[key],datetime) else values[key].isoformat()
                    a=self._save_asset(values)
                    histories=[clean(v) for k,v in row.items() if isinstance(k,str) and re.fullmatch(r'使用人\s*\d+',k) and v]
                    if row.get('完整流转历史'):
                        histories.append(clean(row['完整流转历史'],50000))
                    if row.get('__history'):
                        histories += [str(h.get('at',''))+' '+str(h.get('action',''))+' '+str(h.get('detail','')) for h in row['__history']]
                    for text in histories:
                        self.event(a,'原表历史',text)
                    original=clean(row.get('当前员工编号') or row.get('当前使用人') or '')
                    if original:
                        self.event(a,'原表当前使用信息',original+'；请核对后通过领用操作建立正式员工关联')
                except (ValueError,TypeError,AttributeError) as e:
                    raise BusinessError(f'第 {i} 行：{e}') from e
            self.audit('批量导入',str(len(records))+' 条资产；初始状态为闲置，原使用信息保留为历史')
        return len(records)

    def export_assets(self,path,assets=None):
        self.permit('data.export');self.permit('assets.read')
        allowed={a['id']:a for a in self.list('assets')}
        if assets is not None and any(a['id'] not in allowed for a in assets):raise BusinessError('导出含授权范围外资产')
        assets=list(allowed.values()) if assets is None else [allowed[a['id']] for a in assets]
        headers=list(LABELS.values())
        events=self.list('events'); grouped={}
        for e in sorted(events,key=lambda x:x['created_at']):
            grouped.setdefault(e['asset_id'],[]).append(e['created_at']+' '+e['action']+' '+e['detail'])
        rows=[]
        for a in assets:
            row=[a[k] for k in LABELS]
            row[list(LABELS).index('price_cents')]=str(Decimal(a['price_cents'])/100)
            row[list(LABELS).index('residual_bp')]=str(Decimal(a['residual_bp'])/100)
            rows.append(row+['\n'.join(grouped.get(a['id'],[]))])
        write_table(path,headers+['完整流转历史'],rows)

def safe_cell(v):
    if isinstance(v,str) and v.lstrip().startswith(('=','+','-','@')):
        return "'"+v
    return v

def write_table(path,headers,rows):
    path=Path(path)
    if path.suffix.lower()=='.xlsx':
        from openpyxl import Workbook
        from openpyxl.styles import Font,PatternFill,Alignment
        book=Workbook();sheet=book.active;sheet.title='数据'
        sheet.append(headers)
        for row in rows:
            sheet.append([safe_cell(v) for v in row])
        for c in sheet[1]:
            c.font=Font(color='FFFFFF',bold=True);c.fill=PatternFill('solid',fgColor='365CF5')
        sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
        for col in sheet.columns:
            letter=col[0].column_letter
            sheet.column_dimensions[letter].width=min(50,max(16,max(len(str(c.value or '')) for c in list(col)[:200])+3))
        book.save(path);book.close()
    else:
        with path.open('w',encoding='utf-8-sig',newline='') as f:
            writer=csv.writer(f);writer.writerow(headers);writer.writerows([[safe_cell(v) for v in r] for r in rows])
