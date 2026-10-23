"""2.1 additions: import planning, templates, shared attachments, bulk flows, tasks."""
import base64
import csv
import hashlib
import json
from datetime import date,datetime
from pathlib import Path
from database import COLUMNS
from service import Service as BaseService,BusinessError,clean,integer,money,day,ident,stamp,LABELS,ASSET_FIELDS,write_table

KINDS={'assets':'资产台账','employees':'员工资料','departments':'部门资料','suppliers':'供应商资料'}
CATALOG_LABELS={'departments':{'id':'编号','name':'部门名称','manager':'负责人','notes':'备注'},'employees':{'id':'员工编号','name':'姓名','department_id':'部门编号','phone':'电话','email':'邮箱','active':'启用','notes':'备注'},'suppliers':{'id':'编号','name':'供应商名称','contact':'联系人','phone':'电话','notes':'备注'}}
ATTACHMENT_EXTENSIONS={'.pdf','.docx','.xlsx','.csv','.txt','.png','.jpg','.jpeg','.bmp','.webp','.zip'}

class Service(BaseService):
    def template(self,path,kind='assets',example=True):
        self.require()
        if kind not in KINDS:raise BusinessError('模板类别无效')
        if kind=='assets':
            headers=[LABELS[k] for k in ASSET_FIELDS]+['使用人1','使用人2']
            sample={'id':'A-001','name':'办公笔记本电脑','quantity':1,'category':'笔记本电脑','brand':'ThinkPad','owner':'公司','location':'办公室','price_cents':5000,'life_months':36,'residual_bp':5,'notes':'示例，请替换；单价单位元，残值率单位百分比'}
            rows=[[sample.get(k,'') for k in ASSET_FIELDS]+['','']] if example else []
        else:
            labels=CATALOG_LABELS[kind];headers=list(labels.values())
            sample={'id':{'employees':'E-001','departments':'D-001','suppliers':'S-001'}[kind],'name':{'employees':'示例员工','departments':'示例部门','suppliers':'示例供应商'}[kind],'active':1,'notes':'示例，请替换后导入'}
            rows=[[sample.get(k,'') for k in labels]] if example else []
        write_table(path,headers,rows)

    def sheet_names(self,path):
        from openpyxl import load_workbook
        book=load_workbook(path,read_only=True,data_only=True)
        try:return book.sheetnames
        finally:book.close()

    def read_import(self,path,sheet=None):
        path=Path(path)
        if path.stat().st_size>30*1024*1024:raise BusinessError('导入文件最多 30MB')
        if path.suffix.lower()=='.xlsx':
            from openpyxl import load_workbook
            book=load_workbook(path,read_only=True,data_only=True)
            try:
                ws=book[sheet] if sheet else book.active
                iterator=iter(ws.iter_rows(values_only=True));header=[clean(v) for v in next(iterator,[])]
                while header and not header[-1]:header.pop()
                if not header or len(set(header))!=len(header) or '' in header:raise BusinessError('表头为空或重复，请使用模板')
                records=[]
                for line,row in enumerate(iterator,2):
                    if any(v is not None and str(v).strip() for v in row):
                        records.append(dict(zip(header,row[:len(header)]))|{'__line':line})
            finally:book.close()
        elif path.suffix.lower()=='.csv':
            try:
                with path.open(encoding='utf-8-sig',newline='') as f:
                    reader=csv.DictReader(f)
                    if not reader.fieldnames:raise BusinessError('文件缺少表头')
                    reader.fieldnames=[clean(v) for v in reader.fieldnames]
                    if len(set(reader.fieldnames))!=len(reader.fieldnames) or '' in reader.fieldnames:raise BusinessError('表头为空或重复')
                    records=[]
                    for row in reader:
                        if None in row:raise BusinessError('CSV 某行列数超过表头，请检查逗号和引号')
                        if any(clean(v) for v in row.values()):records.append(row|{'__line':reader.line_num})
            except UnicodeDecodeError:raise BusinessError('CSV 需 UTF-8 编码，请从 Excel 另存为 CSV UTF-8')
        elif path.suffix.lower()=='.json':
            old=json.loads(path.read_text('utf-8-sig'))
            if old.get('version')!=1 or not isinstance(old.get('assets'),list):raise BusinessError('这里只支持旧网页 JSON；桌面备份请使用恢复功能')
            records=[{'编号':a.get('id'),'设备名称':a.get('name'),'数量':a.get('quantity'),'类别':a.get('category'),'品牌型号':a.get('brand'),'归属':a.get('owner'),'其他备注':str(a.get('notes',''))+'\n网页原状态：'+str(a.get('status',''))+'；原使用人：'+str(a.get('user','')),'__history':a.get('history',[]),'__line':i} for i,a in enumerate(old['assets'],2)]
        else:raise BusinessError('支持 Excel .xlsx、CSV UTF-8 和旧网页 JSON')
        if not 1<=len(records)<=50000:raise BusinessError('文件需包含 1～50000 条数据')
        return records

    def _plan(self,records,kind,mode):
        if kind not in KINDS or mode not in ['仅新增','按编号新增或更新']:raise BusinessError('导入选项无效')
        existing={r['id']:r for r in self.db.rows(kind)}
        names={r['name']:r['id'] for r in existing.values()} if kind in ['departments','suppliers'] else {}
        ids=set();rows=[];errors=[];added=updated=0
        for i,row in enumerate(records,2):
            line=row.get('__line',i)
            try:
                if kind=='assets':
                    values={k:row.get(LABELS[k],'') for k in ASSET_FIELDS}
                    values['id']=row.get('编号') or row.get('资产编号') or ''
                    values['name']=row.get('设备名称') or row.get('资产名称') or ''
                    values['category']=row.get('类别') or '未分类'
                    values['brand']=row.get('品牌型号') or row.get('品牌') or ''
                    values['notes']=row.get('其他备注') or row.get('备注') or ''
                    values['quantity']=row.get('数量') if row.get('数量') not in [None,''] else 1
                    values['price_cents']=money(row.get('单价',0));values['life_months']=row.get('使用月数') or 36;values['residual_bp']=money(row.get('残值率',5) if row.get('残值率') not in [None,''] else 5)
                    for k in ['purchase_date','warranty_date']:
                        if isinstance(values[k],(date,datetime)):values[k]=values[k].date().isoformat() if isinstance(values[k],datetime) else values[k].isoformat()
                    data=self.asset_data(values)
                    if data['supplier_id'] and not self.db.one('SELECT id FROM suppliers WHERE id=?',(data['supplier_id'],)):raise BusinessError('供应商编号不存在，请先导入供应商')
                else:
                    labels=CATALOG_LABELS[kind];data={k:clean(row.get(label,'')) for k,label in labels.items()}
                    data['id']=clean(data['id'],100)
                    if not data['id'] or not data['name']:raise BusinessError('编号和名称必填')
                    if kind=='employees':
                        active=clean(row.get('启用',1))
                        if active not in ['0','1']:raise BusinessError('启用需填写 1 或 0')
                        data['active']=int(active)
                        if data['department_id'] and not self.db.one('SELECT id FROM departments WHERE id=?',(data['department_id'],)):raise BusinessError('部门编号不存在，请先导入部门')
                        if not data['active'] and self.db.one('SELECT id FROM assets WHERE employee_id=?',(data['id'],)):raise BusinessError('员工有在用资产，不能停用')
                    else:
                        if data['name'] in names and names[data['name']]!=data['id']:raise BusinessError('部门或供应商名称重复')
                        names[data['name']]=data['id']
                id=data['id']
                if id in ids:raise BusinessError('文件内编号重复：'+id)
                ids.add(id);old=existing.get(id)
                if kind=='assets':
                    if old:self.check_asset(id)
                    c=self.profile()
                    if c['scope']=='self' or (c['scope']=='locations' and data['location'] not in c['locations']):raise BusinessError('导入资产不在可维护的数据范围内')
                    self.permit('assets.edit' if old else 'assets.create')
                if old and mode=='仅新增':raise BusinessError('编号已存在：'+id)
                if kind=='assets' and old and old['status']=='已报废':raise BusinessError('报废资产不能更新')
                histories=[clean(v) for k,v in row.items() if isinstance(k,str) and k.startswith('使用人') and k[3:].strip().isdigit() and v]
                if row.get('完整流转历史'):histories.append(clean(row['完整流转历史'],50000))
                if row.get('__history'):histories.extend(clean(str(h.get('at',''))+' '+str(h.get('action',''))+' '+str(h.get('detail',''))) for h in row['__history'])
                original=clean(row.get('当前使用人') or row.get('当前员工编号') or '')
                if original:histories.append('原表当前使用信息：'+original+'；请核对后正式领用')
                rows.append({'line':line,'data':data,'revision':old.get('revision') if old else None,'existing':bool(old),'previous':old,'history':histories})
                if old:updated+=1
                else:added+=1
            except (ValueError,TypeError,AttributeError) as e:errors.append({'line':line,'id':clean(row.get('编号') or row.get('员工编号') or row.get('资产编号') or ''),'error':str(e)})
        return {'kind':kind,'mode':mode,'rows':rows,'errors':errors,'added':added,'updated':updated,'total':len(records)}

    def plan_import(self,records,kind,mode):
        self.permit('data.import');self.permit('assets.read' if kind=='assets' else kind+'.write')
        if kind!='assets':self.global_scope()
        with self.db.transaction():
            self.require(write=True);return self._plan(records,kind,mode)

    def apply_import(self,records,kind,mode,filename,preview=None):
        self.permit('data.import');self.permit('assets.read' if kind=='assets' else kind+'.write')
        if kind!='assets':self.global_scope()
        with self.db.transaction():
            self.require(write=True);plan=self._plan(records,kind,mode)
            if plan['errors']:raise BusinessError('数据已变化或存在错误，未导入：'+plan['errors'][0]['error'])
            if preview is not None and [(r['data']['id'],r['previous']) for r in preview['rows']]!=[(r['data']['id'],r['previous']) for r in plan['rows']]:raise BusinessError('数据库在预览后已变化，请重新预览再导入')
            for item in plan['rows']:
                if kind=='assets':
                    a=self._save_asset(item['data'],item['revision'])
                    for history in item['history']:self.event(a,'原表历史',history)
                else:self.catalog(kind,item['data'],item['existing'])
            self.db.insert('import_batches',{'id':ident(),'filename':clean(Path(filename).name,300),'kind':KINDS[kind],'mode':mode,'added':plan['added'],'updated':plan['updated'],'actor_id':self.user_id,'created_at':stamp()})
            self.audit('文件导入',KINDS[kind]+' / '+clean(Path(filename).name)+' / 新增 '+str(plan['added'])+' / 更新 '+str(plan['updated']))
        return plan

    def batch_flow(self,selections,action,**kwargs):
        if action not in ['领用','归还','移动位置']:raise BusinessError('批量操作仅支持领用、归还和位置移动')
        if not selections or len({x['id'] for x in selections})!=len(selections):raise BusinessError('请选择不同的资产')
        with self.db.transaction():
            self.require(write=True)
            for a in selections:self.flow(a['id'],a['revision'],action,**kwargs)
            self.audit('批量流转',action+' / '+str(len(selections))+' 条，事务统一提交')

    def attachments(self,asset_id):
        self.permit('attachments.read');self.check_asset(asset_id)
        self.require();return self.db.all('SELECT id,asset_id,filename,size_bytes,sha256,actor_id,created_at FROM attachments WHERE asset_id=? ORDER BY created_at DESC',(asset_id,))

    def upload_attachments(self,asset_id,paths):
        self.permit('attachments.write');self.check_asset(asset_id)
        if not paths:raise BusinessError('请选择文件')
        payload=[]
        for path in paths:
            file=Path(path)
            if file.suffix.lower() not in ATTACHMENT_EXTENSIONS:raise BusinessError('不支持的附件类型：'+file.name)
            if file.stat().st_size>10*1024*1024:raise BusinessError('单个附件最多 10MB：'+file.name)
            raw=file.read_bytes()
            payload.append({'id':ident(),'asset_id':asset_id,'filename':clean(file.name,255),'size_bytes':len(raw),'sha256':hashlib.sha256(raw).hexdigest(),'content_b64':base64.b64encode(raw).decode(),'actor_id':self.user_id,'created_at':stamp()})
        with self.db.transaction():
            self.require(write=True)
            if not self.db.one('SELECT id FROM assets WHERE id=?',(asset_id,)):raise BusinessError('资产不存在')
            total=self.db.one('SELECT COALESCE(SUM(size_bytes),0) AS n FROM attachments')['n']
            if total+sum(p['size_bytes'] for p in payload)>30*1024*1024:raise BusinessError('当前数据库附件总容量最多 30MB，请先下载并移除旧附件')
            for p in payload:self.db.insert('attachments',p)
            self.audit('上传资产附件',asset_id+' / '+'；'.join(p['filename'] for p in payload))
        return len(payload)

    def download_attachment(self,id,path):
        self.permit('attachments.read');self.check_row('attachments',id)
        self.require();row=self.db.one('SELECT * FROM attachments WHERE id=?',(id,))
        if not row:raise BusinessError('附件不存在')
        raw=base64.b64decode(row['content_b64'],validate=True)
        if hashlib.sha256(raw).hexdigest()!=row['sha256']:raise BusinessError('附件内容校验失败')
        Path(path).write_bytes(raw)

    def remove_attachment(self,id):
        self.permit('attachments.write')
        a=self.db.one('SELECT asset_id FROM attachments WHERE id=?',(id,))
        if not a:raise BusinessError('附件不存在')
        self.check_asset(a['asset_id'])
        with self.db.transaction():
            self.require(write=True);a=self.db.one('SELECT * FROM attachments WHERE id=?',(id,))
            if not a:raise BusinessError('附件不存在')
            self.db.execute('DELETE FROM attachments WHERE id=?',(id,));self.audit('移除附件',a['asset_id']+' / '+a['filename'])

    def save_task(self,values):
        self.permit('tasks.write')
        if values.get('asset_id'):self.check_asset(values['asset_id'])
        title=clean(values.get('title'),300);due=day(values.get('due_date',''));priority=values.get('priority','普通');asset_id=clean(values.get('asset_id',''),300)
        if not title or priority not in ['普通','重要','紧急']:raise BusinessError('请填写标题并选择有效优先级')
        with self.db.transaction():
            self.require(write=True)
            if asset_id and not self.db.one('SELECT id FROM assets WHERE id=?',(asset_id,)):raise BusinessError('关联资产不存在')
            id=ident();self.db.insert('tasks',{'id':id,'title':title,'asset_id':asset_id,'priority':priority,'due_date':due,'notes':clean(values.get('notes','')),'status':'待处理','creator_id':self.user_id,'created_at':stamp(),'completed_at':''});self.audit('新增待办',title);return id

    def finish_task(self,id,reopen=False):
        self.permit('tasks.write');self.check_row('tasks',id)
        with self.db.transaction():
            self.require(write=True);task=self.db.one('SELECT * FROM tasks WHERE id=?',(id,))
            if not task:raise BusinessError('待办不存在')
            self.db.update('tasks',id,{'status':'待处理' if reopen else '已完成','completed_at':'' if reopen else stamp()});self.audit('重新开启待办' if reopen else '完成待办',task['title'])

    def update_task(self,values):
        self.permit('tasks.write')
        title=clean(values.get('title'),300);priority=values.get('priority','普通');asset_id=clean(values.get('asset_id',''),300)
        if not title or priority not in ['普通','重要','紧急']:raise BusinessError('请填写标题并选择有效优先级')
        with self.db.transaction():
            self.require(write=True)
            task=self.db.one('SELECT * FROM tasks WHERE id=?',(values.get('id'),))
            if not task:raise BusinessError('待办不存在')
            if task['status']=='已完成':raise BusinessError('已完成的待办请先重新开启，再修改内容')
            if asset_id:self.check_asset(asset_id)
            self.db.update('tasks',task['id'],{'title':title,'asset_id':asset_id,'priority':priority,'due_date':day(values.get('due_date','')),'notes':clean(values.get('notes',''))})
            self.audit('修改待办',title)

    def delete_task(self,id):
        self.permit('tasks.delete');self.check_row('tasks',id)
        with self.db.transaction():
            self.require(write=True)
            task=self.db.one('SELECT * FROM tasks WHERE id=?',(id,))
            if not task:raise BusinessError('待办不存在')
            self.db.execute('DELETE FROM tasks WHERE id=?',(id,))
            self.audit('删除待办',task['title'])

    def update_repair(self,id,values):
        self.permit('repairs.write');self.check_row('repairs',id)
        problem=clean(values.get('problem'),5000)
        if not problem:raise BusinessError('请填写问题描述')
        with self.db.transaction():
            self.require(write=True)
            r=self.db.one('SELECT * FROM repairs WHERE id=?',(id,))
            if not r:raise BusinessError('维修工单不存在')
            fields={'problem':problem,'vendor':clean(values.get('vendor'),300),'result':clean(values.get('result',''))}
            cost=clean(values.get('cost',''),40)
            if cost:fields['cost_cents']=money(cost)
            self.db.update('repairs',id,fields)
            self.audit('修改维修工单',r['asset_id']+' / '+problem)

    def delete_repair(self,id):
        self.permit('repairs.delete');self.check_row('repairs',id)
        with self.db.transaction():
            self.require(write=True)
            r=self.db.one('SELECT * FROM repairs WHERE id=?',(id,))
            if not r:raise BusinessError('维修工单不存在')
            if r['status']=='维修中':raise BusinessError('维修中的工单不能删除，请先办理「维修完成」或「报废」')
            self.db.execute('DELETE FROM repairs WHERE id=?',(id,))
            self.audit('删除维修工单',r['asset_id']+' / '+r['problem'])

    def rename_attachment(self,id,name):
        self.permit('attachments.write')
        with self.db.transaction():
            self.require(write=True)
            a=self.db.one('SELECT * FROM attachments WHERE id=?',(id,))
            if not a:raise BusinessError('附件不存在')
            self.check_asset(a['asset_id'])
            stem=clean(name,200)
            if not stem:raise BusinessError('请填写新的文件名')
            if any(character in stem for character in '\\/:*?"<>|'):raise BusinessError('文件名不能包含 \\ / : * ? " < > | 等符号')
            renamed=stem+Path(a['filename']).suffix
            if renamed==a['filename']:raise BusinessError('文件名没有变化')
            self.db.update('attachments',id,{'filename':renamed})
            self.audit('重命名附件',a['filename']+' → '+renamed)
